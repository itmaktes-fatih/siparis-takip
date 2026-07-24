from openpyxl import load_workbook
import re


def get_week_from_filename(filename):
    """
    Dosya isminden hafta bilgisini alır.
    Örnek:
    20-24 TEMMUZ(W202630).xlsx
    ->
    30. Hafta
    """
    match = re.search(r'W(\d{4})(\d{2})', filename)

    if match:
        return f"{int(match.group(2))}. Hafta"

    return "Bilinmeyen Hafta"


def read_excel(filepath):
    """
    Excel dosyasını okuyup kayıt listesini döndürür.
    """

    workbook = load_workbook(filepath, data_only=True)

    # Şimdilik ilk sayfayı kullanıyoruz.
    sheet = workbook.active

    headers = {}

    # Başlık satırını bul
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=col).value

        if value:
            headers[str(value).strip()] = col

    week = get_week_from_filename(filepath)

    jobs = []

    for row in range(2, sheet.max_row + 1):

        order_no = sheet.cell(row, headers["Sipariş"]).value
        title = sheet.cell(row, headers["Kısa metin"]).value
        workstation = sheet.cell(row, headers["İşlem işyeri"]).value
        date = sheet.cell(row, headers["En erk.bşl.trh."]).value

        jobs.append({

            "week": week,

            "order_no": str(order_no),

            "title": str(title),

            "workstation": str(workstation),

            "date": str(date),

            "status": "Başlanmadı",

            "status_note": ""

        })

    return jobs
