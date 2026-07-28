import os
import re
from pathlib import Path
from openpyxl import load_workbook


# ----------------------------------------------------
# Download klasöründeki en son Excel dosyasını bulur
# ----------------------------------------------------

def get_latest_excel():

    possible_dirs = [
        "/storage/emulated/0/Download",                 # Android
        str(Path.home() / "Downloads"),                 # Windows
    ]

    excel_files = []

    for folder in possible_dirs:

        if not os.path.exists(folder):
            continue

        for file in os.listdir(folder):

            if file.lower().endswith(".xlsx"):

                full_path = os.path.join(folder, file)

                excel_files.append(full_path)

    if not excel_files:
        return None

    excel_files.sort(key=os.path.getmtime, reverse=True)

    return excel_files[0]


# ----------------------------------------------------
# Dosya isminden hafta bilgisini al
# ----------------------------------------------------

def get_week_number(filename):

    match = re.search(r'W\d{4}(\d{2})', filename)

    if match:
        return match.group(1)

    return "00"


# ----------------------------------------------------
# Başlık satırını bul
# ----------------------------------------------------

def find_header_row(sheet):

    for row in range(1, 15):

        values = []

        for cell in sheet[row]:
            values.append(str(cell.value).strip() if cell.value else "")

        if "Sipariş" in values:
            return row

    return 1


# ----------------------------------------------------
# Excel Oku
# ----------------------------------------------------

def read_excel():

    excel_path = get_latest_excel()

    if excel_path is None:
        raise Exception("Download klasöründe Excel bulunamadı.")

    workbook = load_workbook(excel_path, data_only=True)

    sheet = workbook.active

    header_row = find_header_row(sheet)

    headers = {}

    for col in range(1, sheet.max_column + 1):

        value = sheet.cell(header_row, col).value

        if value:
            headers[str(value).strip()] = col

    week = get_week_number(os.path.basename(excel_path))

    jobs = []

    for row in range(header_row + 1, sheet.max_row + 1):

        siparis = sheet.cell(row, headers["Sipariş"]).value

        if siparis is None:
            continue

        job = {

            "week": week,

            "siparis": str(siparis),

            "kisa_metin": str(sheet.cell(row, headers["Kısa metin"]).value),

            "isyeri": str(sheet.cell(row, headers["İşlem işyeri"]).value),

            "tarih": str(sheet.cell(row, headers["En erk.bşl.trh."]).value),

            "durum": "Başlanmadı",

            "aciklama": ""

        }

        jobs.append(job)

    return jobs
